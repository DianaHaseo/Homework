import json
import csv
from pathlib import Path
from openpyxl import load_workbook


def read_transactions(file_path: str) -> list[dict]:
    """Читает транзакции из JSON, CSV или XLSX файла"""
    path = Path(file_path)
    if not path.exists():
        return []

    if path.suffix == ".json":
        return load_transactions_from_json(path)
    elif path.suffix == ".csv":
        return load_transactions_from_csv(path)
    elif path.suffix in [".xlsx", ".xls"]:
        return load_transactions_from_xlsx(path)
    else:
        raise ValueError(f"Неподдерживаемый формат файла: {path.suffix}")


def load_transactions_from_json(path: Path) -> list[dict]:
    """Загружает транзакции из JSON файла"""
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def load_transactions_from_csv(path: Path) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader)


def load_transactions_from_xlsx(path: Path) -> list[dict]:
    wb = load_workbook(filename=path)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    transactions = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        transactions.append(dict(zip(headers, row)))
    return transactions
