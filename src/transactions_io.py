import csv
from typing import Any, Dict, List

import openpyxl


def read_transactions_from_csv(file_path: str) -> List[Dict[str, Any]]:
    """
    Считывает финансовые операции из CSV файла и возвращает список словарей с транзакциями.

    :param file_path: путь к CSV файлу
    :return: список транзакций (каждая транзакция — словарь)
    """
    transactions = []
    with open(file_path, newline='', encoding='utf-8') as csvfile:
        reader = csv.DictReader(csvfile)
        for row in reader:
            transactions.append(dict(row))
    return transactions


def read_transactions_from_excel(file_path: str) -> List[Dict[str, Any]]:
    """
    Считывает финансовые операции из Excel файла и возвращает список словарей с транзакциями.

    :param file_path: путь к файлу Excel
    :return: список транзакций (каждая транзакция — словарь)
    """
    transactions = []
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        transaction = {headers[i]: row[i] for i in range(len(headers))}
        transactions.append(transaction)
    return transactions